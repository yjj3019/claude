#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Fills the '점검결과' (E) column of check_reconstructed.xlsx from an
OCP-HCK-Score.sh .txt report, matched by checklist item number.

Usage: python3 fill_checklist.py <report.txt> <template.xlsx> [output.xlsx]

Parsing logic (ITEM_HDR/RESULT_LINE/parse_items, section-2 pipe parsing) is
ported verbatim from the python heredoc embedded in OCP-HCK-Score.sh
(around lines 542-611) — do not re-derive it.
"""
import re
import sys
from datetime import datetime

import openpyxl

DASH = "--------------------------------------------------------"
SEC_RE = re.compile(r"^■ (\d+)\. (.+)$")
ITEM_HDR = re.compile(r"^\[([0-9]-[0-9](?:-[0-9])?)\]\s+(.+)$")
RESULT_LINE = re.compile(r"^\[결과\]\s+(rc=(\d+)|skip|manual)\s*$")
ITEM_NUM_RE = re.compile(r"^\d+-\d+(-\d+)?$")

STATUS_MAP = {"ok": "정상", "attention": "확인필요", "skip": "건너뜀", "manual": "수동확인"}

SHEET_BY_SECTION = {
    "1": "1.Cluster구성",
    "2": "2.ClusterOperator",
    "3": "3.API연동",
    "4": "4.Network",
    "5": "5.Virtualization",
}


def skip_or_attention(body):
    if "건너뜀" in body:
        return "skip"
    if re.search(r"\berror:|Error from server|unauthorized|command not found|no such", body, re.I):
        return "attention"
    return "ok"


def parse_items(body_lines):
    """Split a section's body into item blocks delimited by DASH/[num] desc/Command/DASH."""
    items = []
    i = 0
    n = len(body_lines)
    while i < n:
        if body_lines[i] == DASH and i + 3 < n and ITEM_HDR.match(body_lines[i + 1]):
            m = ITEM_HDR.match(body_lines[i + 1])
            num, desc = m.group(1), m.group(2)
            j = i + 4
            out = []
            result_tag = None
            while j < n and body_lines[j] != DASH:
                rm = RESULT_LINE.match(body_lines[j])
                if rm and result_tag is None:
                    result_tag = rm.group(1)
                else:
                    out.append(body_lines[j])
                j += 1
            body_text = "\n".join(out).strip("\n")
            if result_tag == "skip":
                status = "skip"
            elif result_tag == "manual":
                status = "manual"
            elif result_tag is not None:
                status = "ok" if result_tag.split("=", 1)[1] == "0" else "attention"
            else:
                status = skip_or_attention(body_text)
            items.append({"num": num, "desc": desc, "status": status})
            i = j
        else:
            i += 1
    return items


def parse_report(text):
    """Returns {item_num: result_text} across all 5 sections."""
    lines = text.split("\n")
    sections = []
    cur = None
    for ln in lines:
        m = SEC_RE.match(ln)
        if m:
            cur = {"num": m.group(1), "lines": []}
            sections.append(cur)
        elif cur is not None and not ln.startswith("════"):
            cur["lines"].append(ln)

    results = {}

    sec2 = next((s for s in sections if s["num"] == "2"), None)
    if sec2:
        for ln in sec2["lines"]:
            if ln.startswith("2-") and "|" in ln:
                parts = [p.strip() for p in ln.split("|")]
                if len(parts) >= 6:
                    results[parts[0]] = parts[5]
                else:
                    results[parts[0]] = "[jq 없음 - 원본 확인]"

    for sec_num in ("1", "3", "4", "5"):
        sec = next((s for s in sections if s["num"] == sec_num), None)
        if not sec:
            continue
        for it in parse_items(sec["lines"]):
            results[it["num"]] = STATUS_MAP.get(it["status"], it["status"])

    return results


def fill_workbook(wb, results):
    """Writes results into each sheet's E column. Returns (filled, missing_in_report)."""
    filled = []
    missing = []
    for sec_num, sheet_name in SHEET_BY_SECTION.items():
        if sheet_name not in wb.sheetnames:
            print("[경고] 시트 없음: %s" % sheet_name, file=sys.stderr)
            continue
        ws = wb[sheet_name]
        for row in ws.iter_rows(min_row=2):
            num_cell = row[0]
            num = num_cell.value
            if not isinstance(num, str) or not ITEM_NUM_RE.match(num):
                continue
            if num in results:
                row[4].value = results[num]
                filled.append(num)
            else:
                missing.append(num)
                print("[경고] 리포트에서 매칭 실패: %s (%s)" % (num, sheet_name), file=sys.stderr)
    return filled, missing


def main():
    if len(sys.argv) not in (3, 4):
        print("usage: python3 fill_checklist.py <report.txt> <template.xlsx> [output.xlsx]",
              file=sys.stderr)
        sys.exit(2)
    report_path, template_path = sys.argv[1], sys.argv[2]
    output_path = sys.argv[3] if len(sys.argv) == 4 else (
        "checklist_filled_%s.xlsx" % datetime.now().strftime("%Y%m%d-%H%M"))

    with open(report_path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    results = parse_report(text)

    wb = openpyxl.load_workbook(template_path)
    filled, missing = fill_workbook(wb, results)
    wb.save(output_path)

    print("채움: %d개, 매칭 실패: %d개 -> %s" % (len(filled), len(missing), output_path))


SAMPLE_REPORT = """점검 대상 클러스터 : https://api.example:6443
점검 수행 계정     : kubeadmin
점검 시작 시각     : Fri Sep 11 09:00:00 KST 2026

════════════════════════════════════════════════════════
■ 1. Cluster 구성 확인
════════════════════════════════════════════════════════

--------------------------------------------------------
[1-1] Node Join 상태 확인
Command : oc get nodes
--------------------------------------------------------
[결과] rc=0
NAME   STATUS   ROLES
node1  Ready    master

--------------------------------------------------------
[1-2] 전체 Node Role 확인
Command : oc get nodes
--------------------------------------------------------
[결과] rc=1
error: some failure

════════════════════════════════════════════════════════
■ 2. Cluster Operator 상태 확인
════════════════════════════════════════════════════════

번호 | Operator | AVAILABLE | PROGRESSING | DEGRADED | 판정
--------------------------------------------------------------
2-1 | authentication | True | False | False | ✅ 정상
2-2 | baremetal | True | True | False | ⚠ 주의(Progressing)

════════════════════════════════════════════════════════
■ 3. API 연동 확인
════════════════════════════════════════════════════════

--------------------------------------------------------
[3-1-1] Namespace 목록 조회
Command : oc get ns
--------------------------------------------------------
[결과] rc=0
NAME   STATUS
default Active

════════════════════════════════════════════════════════
■ 4. Network 확인
════════════════════════════════════════════════════════

--------------------------------------------------------
[4-1] Node 인터페이스 확인
Command : oc debug node/n1
--------------------------------------------------------
[결과] skip
virtctl 미설치로 건너뜀

════════════════════════════════════════════════════════
■ 5. Virtualization 확인
════════════════════════════════════════════════════════

--------------------------------------------------------
[5-3] VM Console 접속 확인
Command : virtctl console vm1
--------------------------------------------------------
[결과] manual
수동 확인 필요

════════════════════════════════════════════════════════
점검 종료 시각 : Fri Sep 11 09:10:00 KST 2026
════════════════════════════════════════════════════════
"""


def _self_check():
    results = parse_report(SAMPLE_REPORT)
    assert results["1-1"] == "정상", results
    assert results["1-2"] == "확인필요", results
    assert results["2-1"] == "✅ 정상", results
    assert results["2-2"] == "⚠ 주의(Progressing)", results
    assert results["3-1-1"] == "정상", results
    assert results["4-1"] == "건너뜀", results
    assert results["5-3"] == "수동확인", results
    assert "1-3" not in results

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sec_num, sheet_name in SHEET_BY_SECTION.items():
        ws = wb.create_sheet(sheet_name)
        ws.append(["번호", "점검 항목", "Command", "정상 판정 기준", "점검결과", "비고"])
    ws1 = wb["1.Cluster구성"]
    ws1.append(["1-1", "d", "c", "s", None, None])
    ws1.append(["1-2", "d", "c", "s", None, None])
    ws1.append(["1-3", "d", "c", "s", None, None])
    ws3 = wb["3.API연동"]
    ws3.append(["3-1. 배너 행", None, None, None, None, None])
    ws3.append(["3-1-1", "d", "c", "s", None, None])
    ws2 = wb["2.ClusterOperator"]
    ws2.append(["2-10", "d", "c", "s", None, None])  # regression: two-digit numbers (2-10..2-32) must still match

    results["2-10"] = "✅ 정상"
    filled, missing = fill_workbook(wb, results)
    assert ws1["E2"].value == "정상"
    assert ws1["E3"].value == "확인필요"
    assert "1-3" in missing
    assert ws3["A2"].value.startswith("3-1.") and ws3["E2"].value is None
    assert ws3["E3"].value == "정상"
    assert ws2["E2"].value == "✅ 정상", "two-digit item number (2-10) not matched by ITEM_NUM_RE"
    print("self-check OK: filled=%d missing=%s" % (len(filled), missing))


if __name__ == "__main__":
    if len(sys.argv) == 2 and sys.argv[1] == "--self-check":
        _self_check()
    else:
        main()
